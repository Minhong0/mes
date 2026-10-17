"""부트캠프 자료를 MSSQL seed SQL로 변환한다.

입력
  data/master/*.csv              기준정보 (사람이 관리하는 원본)
  data/raw/(2차 업무자료)*.xlsx    2025년 10월 일별 생산실적, 비가동 내역, OEE 정답

출력
  db/seed/01_master.sql          기준정보 INSERT
  db/seed/02_history.sql         부트캠프 과거 실적 INSERT (production_daily, downtime)
  data/profile/downtime_events.csv  설비 시뮬레이터가 재현할 실제 정지 기록 (날짜, 설비, 사유 코드, 정지 시간)

실행
  python tools/import_excel.py

같은 입력이면 항상 같은 출력이 나오도록 정렬하고, 생성 시각 같은 값은 넣지 않는다.
변환 전에 데이터를 검증하고, 하나라도 실패하면 파일을 쓰지 않고 종료한다.
"""

from __future__ import annotations

import csv
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MASTER_DIR = ROOT / "data" / "master"
RAW_DIR = ROOT / "data" / "raw"
SEED_DIR = ROOT / "db" / "seed"
PROFILE_DIR = ROOT / "data" / "profile"

# MSSQL은 INSERT ... VALUES 한 문장에 최대 1000행까지 허용한다.
MAX_ROWS_PER_INSERT = 1000

# OEE 정답 비교 허용 오차 (%p). 엑셀은 소수점 아래 전체를 계산하고 DB 뷰는 둘째 자리에서 반올림한다.
OEE_TOLERANCE = Decimal("0.01")

CODE_GROUP_NAMES = {
    "ITEM_TYPE": "품목유형",
    "WO_STATUS": "작업지시 상태",
    "EQUIP_STATUS": "설비 상태",
    "DT_CAT": "비가동 대분류",
    "DT_SUB": "비가동 소분류",
    "DEFECT_TYPE": "불량 유형",
}

# CSV 파일 → 테이블. 외래키 순서대로 INSERT 해야 하므로 순서가 중요하다.
MASTER_TABLES = [
    ("common_codes", "common_code"),
    ("items", "item"),
    ("processes", "process"),
    ("bom", "bom"),
    ("work_centers", "work_center"),
    ("routings", "routing"),
    ("equipment", "equipment"),
    ("inspection_specs", "inspection_spec"),
    ("workers", "worker"),
]

# CSV에서 숫자로 넣을 컬럼. 나머지는 문자열로 넣는다.
NUMERIC_COLUMNS = {
    "qty", "std_time_min", "batch_max_qty", "setup_time_min", "workers", "install_year",
    "ideal_ct_min", "planned_min_per_day", "nominal", "lsl", "usl", "sort_order",
}

# 한글이 들어갈 수 있는 컬럼은 N'...'(NVARCHAR)로 넣는다.
UNICODE_COLUMNS = {
    "code_name", "group_name", "item_name", "spec", "note", "process_name", "department",
    "wc_name", "location", "equip_name", "inspect_name", "worker_name",
}


class ValidationError(Exception):
    pass


# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------

def read_csv(name: str) -> list[dict[str, str]]:
    with open(MASTER_DIR / f"{name}.csv", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def sql_value(column: str, value) -> str:
    """파이썬 값을 SQL 리터럴로 바꾼다. 빈 값은 NULL."""
    if value is None or value == "":
        return "NULL"
    if column in NUMERIC_COLUMNS or isinstance(value, (int, Decimal)):
        return str(value)
    text = str(value).replace("'", "''")   # 작은따옴표 이스케이프
    prefix = "N" if column in UNICODE_COLUMNS else ""
    return f"{prefix}'{text}'"


def insert_statements(table: str, columns: list[str], rows: list[dict]) -> list[str]:
    """행 목록을 1000행 단위 INSERT 문으로 나눈다."""
    statements = []
    for start in range(0, len(rows), MAX_ROWS_PER_INSERT):
        chunk = rows[start:start + MAX_ROWS_PER_INSERT]
        values = ",\n".join(
            "    (" + ", ".join(sql_value(c, r.get(c)) for c in columns) + ")" for r in chunk
        )
        statements.append(f"INSERT INTO {table} ({', '.join(columns)}) VALUES\n{values};")
    return statements


def file_header(title: str) -> str:
    return (
        f"-- {title}\n"
        "-- 자동 생성 파일입니다. 직접 수정하지 말고 tools/import_excel.py를 다시 실행하세요.\n"
        "USE minimes;\nGO\nSET ANSI_NULLS ON;\nSET QUOTED_IDENTIFIER ON;\nSET NOCOUNT ON;\nGO\n\n"
    )


def find_raw(prefix: str) -> Path:
    matches = sorted(RAW_DIR.glob(f"{prefix}*.xlsx"))
    if len(matches) != 1:
        raise ValidationError(f"data/raw에서 '{prefix}'로 시작하는 엑셀을 1개 찾아야 하는데 {len(matches)}개입니다.")
    return matches[0]


def to_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def to_time(value) -> time:
    if isinstance(value, time):
        return value
    return datetime.strptime(str(value).strip(), "%H:%M").time()


# ---------------------------------------------------------------------------
# 1. 기준정보
# ---------------------------------------------------------------------------

def load_master() -> dict[str, list[dict[str, str]]]:
    return {name: read_csv(name) for name, _ in MASTER_TABLES}


def validate_master(m: dict[str, list[dict[str, str]]]) -> None:
    """CSV끼리 참조하는 코드가 모두 존재하는지 확인한다 (DB의 외래키를 미리 검사)."""
    items = {r["item_code"] for r in m["items"]}
    processes = {r["process_code"] for r in m["processes"]}
    wcs = {r["wc_code"] for r in m["work_centers"]}
    equipment = {r["equip_code"] for r in m["equipment"]}
    errors = []

    def check(where: str, value: str, valid: set[str], allow_empty: bool = False) -> None:
        if allow_empty and value == "":
            return
        if value not in valid:
            errors.append(f"{where}: '{value}' 이(가) 기준정보에 없습니다.")

    for r in m["bom"]:
        check("bom.parent_item_code", r["parent_item_code"], items)
        check("bom.child_item_code", r["child_item_code"], items)
        check("bom.process_code", r["process_code"], processes)
    for r in m["work_centers"]:
        check("work_centers.process_code", r["process_code"], processes)
    for r in m["routings"]:
        check("routings.item_code", r["item_code"], items)
        check("routings.process_code", r["process_code"], processes)
        check("routings.wc_code", r["wc_code"], wcs)
        check("routings.output_item_code", r["output_item_code"], items)
    for r in m["equipment"]:
        check("equipment.wc_code", r["wc_code"], wcs)
    for r in m["inspection_specs"]:
        check("inspection_specs.process_code", r["process_code"], processes)
        check("inspection_specs.equip_code", r["equip_code"], equipment, allow_empty=True)
    for r in m["common_codes"]:
        if r["group_code"] not in CODE_GROUP_NAMES:
            errors.append(f"common_codes: 그룹 '{r['group_code']}'의 이름이 CODE_GROUP_NAMES에 없습니다.")

    if errors:
        raise ValidationError("기준정보 검증 실패\n  " + "\n  ".join(errors))


def build_master_sql(m: dict[str, list[dict[str, str]]]) -> str:
    parts = [file_header("기준정보 초기 데이터 (출처: data/master/*.csv)")]

    groups = [{"group_code": g, "group_name": n} for g, n in CODE_GROUP_NAMES.items()]
    parts += insert_statements("common_code_group", ["group_code", "group_name"], groups)

    for csv_name, table in MASTER_TABLES:
        rows = m[csv_name]
        columns = list(rows[0].keys())
        parts.append(f"\n-- {table}: {len(rows)}건")
        parts += insert_statements(table, columns, rows)

    return "\n".join(parts) + "\nGO\n"


# ---------------------------------------------------------------------------
# 2. 부트캠프 과거 실적 (2차 자료)
# ---------------------------------------------------------------------------

@dataclass
class DailyRow:
    prod_date: date
    equip_code: str
    planned_min: int
    downtime_min: int
    run_min: int
    prod_qty: int
    good_qty: int
    defect_qty: int
    ideal_ct_min: Decimal


@dataclass
class DowntimeRow:
    equip_code: str
    started_at: datetime
    ended_at: datetime
    duration_min: int
    dt_cat: str
    dt_sub: str


def sheet_rows(ws, header_first_cell: str) -> list[tuple]:
    """헤더 행(B열 값이 header_first_cell인 행) 아래의 데이터 행을 B열부터 읽는다."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows):
        if row[1] == header_first_cell:
            return [r[1:] for r in rows[i + 1:] if r[1] is not None]
    raise ValidationError(f"시트 '{ws.title}'에서 '{header_first_cell}' 헤더를 찾지 못했습니다.")


def load_daily(wb) -> list[DailyRow]:
    result = []
    for r in sheet_rows(wb["일별 생산실적"], "날짜"):
        d, code, _name, planned, down, run, prod, good, defect, ct = r[:10]
        result.append(DailyRow(to_date(d), code, int(planned), int(down), int(run),
                               int(prod), int(good), int(defect), Decimal(str(ct))))
    return result


def load_downtime(wb, codes: list[dict[str, str]]) -> list[DowntimeRow]:
    cat_by_name = {c["code_name"]: c["code"] for c in codes if c["group_code"] == "DT_CAT"}
    sub_by_name = {(c["parent_code"], c["code_name"]): c["code"] for c in codes if c["group_code"] == "DT_SUB"}
    result, errors = [], []
    for r in sheet_rows(wb["비가동 내역"], "날짜"):
        d, code, _name, start, end, minutes, cat_name, sub_name = r[:8]
        cat = cat_by_name.get(cat_name)
        sub = sub_by_name.get((cat, sub_name))
        if cat is None or sub is None:
            errors.append(f"{d} {code}: 사유 '{cat_name} / {sub_name}'에 맞는 코드가 없습니다.")
            continue
        day = to_date(d)
        started = datetime.combine(day, to_time(start))
        ended = datetime.combine(day, to_time(end))
        if ended < started:              # 자정을 넘긴 정지 (야간 교대)
            ended += timedelta(days=1)
        result.append(DowntimeRow(code, started, ended, int(minutes), cat, sub))
    if errors:
        raise ValidationError("비가동 사유 코드 변환 실패\n  " + "\n  ".join(errors))
    return result


def load_expected_oee(wb) -> dict[str, Decimal]:
    """'OEE 산출 과정' 시트 STEP 4의 설비별 월간 OEE(정답)를 읽는다."""
    expected = {}
    for row in wb["OEE 산출 과정"].iter_rows(values_only=True):
        code, oee = row[1], row[11] if len(row) > 11 else None
        if isinstance(code, str) and code.startswith("CNC") and isinstance(oee, (int, float)):
            expected[code] = Decimal(str(oee))
    if not expected:
        raise ValidationError("'OEE 산출 과정' 시트에서 설비별 OEE 정답을 찾지 못했습니다.")
    return expected


def validate_history(daily: list[DailyRow], downtime: list[DowntimeRow], equipment: set[str]) -> list[str]:
    """DB 제약조건과 같은 규칙으로 검사한다. 오류는 예외로, 참고할 점은 경고 목록으로 돌려준다."""
    errors, warnings = [], []
    seen = set()
    for r in daily:
        key = (r.prod_date, r.equip_code)
        if key in seen:
            errors.append(f"일별 실적 중복: {key}")
        seen.add(key)
        if r.equip_code not in equipment:
            errors.append(f"일별 실적 {key}: 설비 '{r.equip_code}'가 기준정보에 없습니다.")
        if r.run_min != r.planned_min - r.downtime_min:
            errors.append(f"일별 실적 {key}: 실가동({r.run_min}) ≠ 계획({r.planned_min}) - 비가동({r.downtime_min})")
        if r.prod_qty != r.good_qty + r.defect_qty:
            errors.append(f"일별 실적 {key}: 생산({r.prod_qty}) ≠ 양품({r.good_qty}) + 불량({r.defect_qty})")
    for r in downtime:
        if r.equip_code not in equipment:
            errors.append(f"비가동 {r.started_at}: 설비 '{r.equip_code}'가 기준정보에 없습니다.")
        span = int((r.ended_at - r.started_at).total_seconds() // 60)
        if span != r.duration_min:
            errors.append(f"비가동 {r.equip_code} {r.started_at}: 시작~종료 {span}분 ≠ 정지시간 {r.duration_min}분")
    if errors:
        raise ValidationError("과거 실적 검증 실패\n  " + "\n  ".join(errors))

    # 비가동 상세 합계와 일별 실적의 비가동시간이 다른 날 (원본 자료의 특성, 오류로 보지 않음)
    detail = defaultdict(int)
    for r in downtime:
        detail[(r.started_at.date(), r.equip_code)] += r.duration_min
    diff = [(k, v.downtime_min, detail.get(k, 0)) for v in daily for k in [(v.prod_date, v.equip_code)]
            if detail.get(k, 0) != v.downtime_min]
    if diff:
        warnings.append(f"일별 실적의 비가동시간과 비가동 상세 합계가 다른 설비·일자: {len(diff)}/{len(daily)}건 "
                        f"(예: {diff[0][0][0]} {diff[0][0][1]} 일별 {diff[0][1]}분, 상세 {diff[0][2]}분)")
    days = sorted({r.prod_date for r in daily})
    warnings.append(f"일별 실적 기간: {days[0]} ~ {days[-1]}, {len(days)}일")
    return warnings


def monthly_oee(daily: list[DailyRow]) -> dict[str, Decimal]:
    """v_oee_monthly 뷰와 같은 방식: 시간과 수량을 먼저 합산한 뒤 OEE를 계산한다."""
    total = defaultdict(lambda: defaultdict(Decimal))
    for r in daily:
        t = total[r.equip_code]
        t["planned"] += r.planned_min
        t["run"] += r.run_min
        t["prod"] += r.prod_qty
        t["good"] += r.good_qty
        t["ideal"] += r.prod_qty * r.ideal_ct_min
    return {
        code: 100 * (t["run"] / t["planned"]) * (t["ideal"] / t["run"]) * (t["good"] / t["prod"])
        for code, t in total.items()
    }


def check_oee(daily: list[DailyRow], expected: dict[str, Decimal]) -> list[str]:
    actual = monthly_oee(daily)
    lines, failed = [], False
    for code in sorted(expected):
        got = actual.get(code)
        ok = got is not None and abs(got - expected[code]) <= OEE_TOLERANCE
        failed |= not ok
        lines.append(f"  {code}: 계산 {got:.2f}%  정답 {expected[code]:.2f}%  {'OK' if ok else 'NG'}")
    if failed:
        raise ValidationError("월간 OEE가 엑셀 정답과 다릅니다\n" + "\n".join(lines))
    return lines


def build_history_sql(daily: list[DailyRow], downtime: list[DowntimeRow]) -> str:
    parts = [file_header("부트캠프 과거 실적 (출처: 2차 업무자료, 2025년 10월)")]
    parts.append("-- 다시 실행해도 결과가 같도록 기존 부트캠프 데이터를 지우고 넣는다.")
    parts.append("DELETE FROM downtime WHERE source = 'BOOTCAMP';")
    parts.append("DELETE FROM production_daily WHERE source = 'BOOTCAMP';\n")

    daily_rows = [
        {"prod_date": r.prod_date.isoformat(), "equip_code": r.equip_code, "planned_min": r.planned_min,
         "downtime_min": r.downtime_min, "prod_qty": r.prod_qty, "good_qty": r.good_qty,
         "defect_qty": r.defect_qty, "ideal_ct_min": r.ideal_ct_min, "source": "BOOTCAMP"}
        for r in sorted(daily, key=lambda x: (x.prod_date, x.equip_code))
    ]
    parts.append(f"-- production_daily: {len(daily_rows)}건 (run_min은 계산 컬럼이라 넣지 않음)")
    parts += insert_statements("production_daily", list(daily_rows[0].keys()), daily_rows)

    dt_rows = [
        {"equip_code": r.equip_code, "started_at": r.started_at.isoformat(sep=" "),
         "ended_at": r.ended_at.isoformat(sep=" "), "duration_min": r.duration_min,
         "dt_cat": r.dt_cat, "dt_sub": r.dt_sub, "source": "BOOTCAMP"}
        for r in sorted(downtime, key=lambda x: (x.started_at, x.equip_code))
    ]
    parts.append(f"\n-- downtime: {len(dt_rows)}건")
    parts += insert_statements("downtime", list(dt_rows[0].keys()), dt_rows)

    return "\n".join(parts) + "\nGO\n"


# ---------------------------------------------------------------------------

def build_downtime_events_csv(downtime: list[DowntimeRow]) -> str:
    """시뮬레이터용 정지 기록. 원본의 한글 사유 대신 코드만 담고, 거래처 등 다른 정보는 넣지 않는다."""
    lines = ["prod_date,equip_code,dt_cat,dt_sub,duration_min"]
    for r in sorted(downtime, key=lambda x: (x.started_at, x.equip_code)):
        lines.append(f"{r.started_at.date().isoformat()},{r.equip_code},{r.dt_cat},{r.dt_sub},{r.duration_min}")
    return "\n".join(lines) + "\n"


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")
    print(f"작성: {path.relative_to(ROOT)}")


def main() -> int:
    try:
        master = load_master()
        validate_master(master)
        print(f"기준정보 검증 통과: " + ", ".join(f"{n} {len(master[n])}건" for n, _ in MASTER_TABLES))

        wb = openpyxl.load_workbook(find_raw("(2차"), data_only=True, read_only=True)
        daily = load_daily(wb)
        downtime = load_downtime(wb, master["common_codes"])
        equipment = {r["equip_code"] for r in master["equipment"]}
        warnings = validate_history(daily, downtime, equipment)
        print(f"과거 실적 검증 통과: 일별 실적 {len(daily)}건, 비가동 {len(downtime)}건")
        for w in warnings:
            print(f"  참고: {w}")

        print("월간 OEE 정답 비교:")
        for line in check_oee(daily, load_expected_oee(wb)):
            print(line)
    except ValidationError as e:
        print(f"\n[중단] {e}", file=sys.stderr)
        return 1

    write(SEED_DIR / "01_master.sql", build_master_sql(master))
    write(SEED_DIR / "02_history.sql", build_history_sql(daily, downtime))
    write(PROFILE_DIR / "downtime_events.csv", build_downtime_events_csv(downtime))
    return 0


if __name__ == "__main__":
    sys.exit(main())
